from PyQt5.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton, QTableView, QTabWidget, QFileDialog, QMessageBox, QHeaderView, QButtonGroup, QSizePolicy)
from PyQt5.QtCore import Qt, QAbstractTableModel, QVariant
from input_validation import parse_integer
import openpyxl
from openpyxl.styles import Font
from technical_scoring import NoteTech
from database import Database
from docx import Document
from word_export import configure_landscape_document, format_landscape_table
import os


class DataTableModel(QAbstractTableModel):
    def __init__(self, data, headers, parent=None):
        super(DataTableModel, self).__init__(parent)
        self._data = data
        self._headers = headers

    def rowCount(self, parent=None):
        return len(self._data)

    def columnCount(self, parent=None):
        return len(self._headers)

    def data(self, index, role=Qt.DisplayRole):
        if index.isValid():
            if role == Qt.DisplayRole:
                value = self._data[index.row()][index.column()]
                # Convertir la valeur en chaîne de caractères ou autre type simple
                return str(value) if value is not None else ""
        return QVariant()

    def headerData(self, section, orientation, role=Qt.DisplayRole):
        if role == Qt.DisplayRole:
            if orientation == Qt.Horizontal:
                return self._headers[section]
        return QVariant()


class TechnicalEvaluation(QWidget):
    def __init__(self, db):
        super().__init__()
        self.db = db
        self.initUI()

    def initUI(self):
        # Layout principal
        layout = QVBoxLayout()
        self.setLayout(layout)

        lot_selector_layout = QHBoxLayout()
        lot_selector_layout.setSpacing(6)
        layout.addLayout(lot_selector_layout)

        lot_selector_layout.addWidget(QLabel("Lot :"))
        self.lot_button_group = QButtonGroup(self)
        self.lot_button_group.setExclusive(True)
        self.lot_buttons = []
        self.selected_lot = 1
        for lot_number in range(1, 12):
            button = QPushButton(f"Lot {lot_number}")
            button.setObjectName("lot_selector_button")
            button.setCheckable(True)
            button.setMinimumHeight(36)
            button.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
            button.clicked.connect(
                lambda checked=False, number=lot_number: self.select_lot(number)
            )
            self.lot_button_group.addButton(button, lot_number)
            self.lot_buttons.append(button)
            lot_selector_layout.addWidget(button, 1)
        self.lot_buttons[0].setChecked(True)

        self.table_tabs = QTabWidget()
        self.summary_table = QTableView()
        self.summary_table.setAlternatingRowColors(True)
        self.summary_table.doubleClicked.connect(self.show_detail_row)
        self.table_tabs.addTab(self.summary_table, "Synthèse")

        self.table = QTableView()
        self.table.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.table_tabs.addTab(self.table, "Critères détaillés")
        layout.addWidget(self.table_tabs)

        # Boutons d'export
        hbox2 = QHBoxLayout()
        layout.addLayout(hbox2)

        self.export_button_excel = QPushButton("Exporter vers Excel")
        self.export_button_excel.clicked.connect(self.export_to_excel)
        hbox2.addWidget(self.export_button_excel)

        self.export_button_word = QPushButton("Exporter vers Word")
        self.export_button_word.clicked.connect(self.export_to_word)
        hbox2.addWidget(self.export_button_word)

        self.load_data(show_empty_message=False)

    def select_lot(self, lot_number):
        self.selected_lot = lot_number
        self.load_data()

    def load_data(self, show_empty_message=True):
        lot_num = self.selected_lot
        self.table.setModel(None)
        self.summary_table.setModel(None)

        try:
            soumissionnaires = self.db.get_soumissionnaires_by_lot(lot_num)

            if not soumissionnaires:
                if show_empty_message:
                    QMessageBox.information(self, "Information", f"Aucun soumissionnaire trouvé pour le lot {lot_num}.")
                return

            note_tech = NoteTech()

            soumissionnaires.sort(
                key=lambda row: (
                    parse_integer(row[0]) is None,
                    parse_integer(row[0]) if parse_integer(row[0]) is not None else str(row[0]),
                )
            )

            headers = [
                'Numéro Pli', 'Nom',
                'Ingénieur', 'Technicien', 'Ouvrier', 'Métreur',
                'Total C1', 'Camion', 'Nacelle', 'Moyen', 'Outillage', 'Total C2',
                'Total C3', 'Total C4', 'Note Tech', 'Statut'
            ]

            data = []

            for (n_pli, nom, lot, ing, ts, ouvrier, metreur, camion, nacelle, moyen_levage, outillage, delai_execution, delai_garantie) in soumissionnaires:
                delai_execution = parse_integer(str(delai_execution).strip(), minimum=0)
                delai_garantie = parse_integer(str(delai_garantie).strip(), minimum=0)
                if delai_execution is None or delai_garantie is None:
                    QMessageBox.warning(self, "Erreur", f"Les délais pour le soumissionnaire {n_pli} ne sont pas valides.")
                    continue

                notes = note_tech.calcule_note_tech(
                    ing, ts, ouvrier, metreur, camion, nacelle, moyen_levage, outillage, delai_execution, delai_garantie, lot
                )

                data.append([
                    n_pli, nom,
                    notes['note_ing'], notes['note_ts'], notes['note_ouvrier'], notes['note_metreur'],
                    notes['note_1'], notes['note_camion'], notes['note_nacelle'], notes['note_moyen_de_levage'],
                    notes['note_outillage'], notes['note_2'],
                    notes['note_3'], notes['note_4'],
                    notes['note_tech'], notes['statut']
                ])

            model = DataTableModel(data, headers)
            self.table.setModel(model)
            summary_data = self.build_summary_rows(data)
            self.summary_table.setModel(
                DataTableModel(
                    summary_data,
                    ["Numéro pli", "Soumissionnaire", "Note technique", "Statut"],
                )
            )
            summary_header = self.summary_table.horizontalHeader()
            summary_header.setSectionResizeMode(0, QHeaderView.ResizeToContents)
            summary_header.setSectionResizeMode(1, QHeaderView.Stretch)
            summary_header.setSectionResizeMode(2, QHeaderView.ResizeToContents)
            summary_header.setSectionResizeMode(3, QHeaderView.ResizeToContents)

        except ValueError:
            QMessageBox.warning(self, "Erreur", "Le numéro de lot doit être un entier valide.")

    @staticmethod
    def build_summary_rows(data):
        return [[row[0], row[1], row[14], row[15]] for row in data]

    def show_detail_row(self, index):
        self.table_tabs.setCurrentWidget(self.table)
        self.table.selectRow(index.row())
        self.table.scrollTo(self.table.model().index(index.row(), 0))


    def export_to_excel(self):
        if not self.table.model():
            QMessageBox.warning(self, "Attention", "Rien à exporter. Veuillez charger les données d'abord.")
            return

        lot_num = str(self.selected_lot)

        # Chemin du bureau de l'utilisateur
        desktop_path = os.path.join(os.path.expanduser("~"), "Desktop")
        file_path = os.path.join(desktop_path, f"Evaluation_Tech_Lot_{lot_num}.xlsx")

        # Vérifier si le fichier existe déjà
        if os.path.exists(file_path):
            reply = QMessageBox.question(self, "Fichier existant", f"Le fichier '{file_path}' existe déjà. Voulez-vous le remplacer ?", QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
            if reply == QMessageBox.No:
                return

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = f"lot{lot_num}"

        headers = [self.table.model().headerData(i, Qt.Horizontal, Qt.DisplayRole) for i in range(self.table.model().columnCount())]
        for col_num, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col_num, value=header)
            cell.font = Font(bold=True)

        for row_num in range(self.table.model().rowCount()):
            for col_num in range(self.table.model().columnCount()):
                value = self.table.model().data(self.table.model().index(row_num, col_num), Qt.DisplayRole)
                sheet.cell(row=row_num + 2, column=col_num + 1, value=value)

        workbook.save(file_path)

        QMessageBox.information(self, "Succès", f"Les données ont été exportées en Excel avec succès : {file_path}")

    def export_to_word(self):
        if not self.table.model():
            QMessageBox.warning(self, "Attention", "Rien à exporter. Veuillez charger les données d'abord.")
            return

        lot_num = str(self.selected_lot)

        # Chemin du bureau de l'utilisateur
        desktop_path = os.path.join(os.path.expanduser("~"), "Desktop")
        file_path, _ = QFileDialog.getSaveFileName(self, "Sauvegarder le fichier Word", os.path.join(desktop_path, f"Evaluation_Tech_Lot_{lot_num}.docx"), "Documents Word (*.docx)")

        if not file_path:
            return

        # Création du document Word
        doc = Document()
        usable_width = configure_landscape_document(doc)

        # Ajout du titre
        doc.add_heading(f"Évaluation Technique - Lot {lot_num}", level=1)

        # Création du tableau
        num_columns = self.table.model().columnCount()
        table = doc.add_table(rows=1, cols=num_columns)

        # Ajout des en-têtes de colonnes
        hdr_cells = table.rows[0].cells
        headers = [self.table.model().headerData(i, Qt.Horizontal, Qt.DisplayRole) for i in range(num_columns)]
        for i, header in enumerate(headers):
            hdr_cells[i].text = str(header)

        # Ajout des données
        for row in range(self.table.model().rowCount()):
            row_cells = table.add_row().cells
            for col in range(num_columns):
                index = self.table.model().index(row, col)
                value = self.table.model().data(index, Qt.DisplayRole)
                row_cells[col].text = str(value) if value is not None else ""

        format_landscape_table(
            table,
            (0.8, 1.8, 0.75, 0.75, 0.7, 0.7, 0.8, 0.7,
             0.7, 0.85, 0.8, 0.8, 0.8, 0.8, 1.0, 1.2),
            usable_width,
            font_size=7.5,
        )

        # Sauvegarde du document
        try:
            doc.save(file_path)
            QMessageBox.information(self, "Succès", f"Les données ont été exportées en Word avec succès : {file_path}")
        except Exception as e:
            QMessageBox.critical(self, "Erreur", f"Une erreur est survenue lors de l'exportation : {e}")



