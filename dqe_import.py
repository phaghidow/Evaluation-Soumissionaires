import hashlib
import re

from openpyxl import load_workbook


def inspect_dqe_workbook(file_path):
    """Analyse un classeur DQE sans modifier la base de données."""
    articles_by_lot = {}
    invalid_rows = []
    checksum = hashlib.sha256()
    with open(file_path, 'rb') as source:
        for block in iter(lambda: source.read(1024 * 1024), b''):
            checksum.update(block)

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        for sheet_name in workbook.sheetnames:
            match = re.search(r"(?:LOT|Lot)\s*(\d+)", str(sheet_name))
            if not match:
                continue
            lot_num = int(match.group(1))
            articles = []
            for row_number, row in enumerate(
                workbook[sheet_name].iter_rows(min_row=2, values_only=True), 2
            ):
                number, quantity = row[0], row[1]
                if number is None or quantity is None:
                    continue
                try:
                    number = int(number)
                    quantity = float(quantity)
                except (TypeError, ValueError):
                    invalid_rows.append(f'{sheet_name}!A{row_number}: valeur invalide')
                    continue
                if number > 0 and quantity >= 0:
                    articles.append((number, quantity))
                else:
                    invalid_rows.append(f'{sheet_name}!A{row_number}: valeur hors limites')
            if articles:
                articles_by_lot[lot_num] = articles
    finally:
        workbook.close()

    if not articles_by_lot:
        raise ValueError("Aucune feuille LOT contenant des articles n'a été trouvée.")

    return {
        'articles_by_lot': articles_by_lot,
        'invalid_rows': invalid_rows,
        'checksum': checksum.hexdigest(),
    }


def import_dqe_workbook(database, file_path, return_details=False):
    """Importe les feuilles LOT du classeur dans la table articles."""
    inspection = inspect_dqe_workbook(file_path)
    articles_by_lot = inspection['articles_by_lot']

    database.backup_database()
    database.replace_articles_from_dqe(articles_by_lot)
    database.record_dqe_import(
        inspection['checksum'],
        sum(len(items) for items in articles_by_lot.values()),
        len(articles_by_lot),
        len(inspection['invalid_rows']),
    )
    result = {
        'total_articles': sum(len(items) for items in articles_by_lot.values()),
        'lot_count': len(articles_by_lot),
        'invalid_rows': inspection['invalid_rows'],
        'checksum': inspection['checksum'],
    }
    return result if return_details else (result['total_articles'], result['lot_count'])
