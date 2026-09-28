import os
import sqlite3
import tempfile
import unittest

from dqe_import import import_dqe_workbook

from database import Database
from input_validation import is_digit_string, parse_amount, parse_integer, parse_integer_tokens
from technical_evaluation import TechnicalEvaluation
from financial.price_schedule import PriceSchedule


class WorkflowDatabaseTests(unittest.TestCase):
    def setUp(self):
        self.database_path = tempfile.mktemp(suffix=".db")
        self.database = Database(self.database_path)

    def tearDown(self):
        self.database.close()
        if os.path.exists(self.database_path):
            os.remove(self.database_path)
        if os.path.exists(f"{self.database_path}.bak"):
            os.remove(f"{self.database_path}.bak")

    def test_dqe_articles_are_replaced_atomically(self):
        self.database.replace_articles_from_dqe({1: [(1, 10), (2, 5.5)]})
        self.assertEqual(self.database.fetch_articles_by_lot(1), [(1, 10.0), (2, 5.5)])

        self.database.replace_articles_from_dqe({1: [(1, 20)]})
        self.assertEqual(self.database.fetch_articles_by_lot(1), [(1, 20.0)])

    def test_dqe_lot_save_is_atomic_and_calculates_totals(self):
        self.database.insert_soumissionnaire("91", "Test DQE", "1")
        self.database.add_offre_qualifie("91", 1, 85)
        self.database.replace_articles_from_dqe({1: [(1, 2), (2, 3)]})

        self.database.save_dqe_lot(
            "91", 1, [(1, 10, 20), (2, 5, 15)]
        )

        self.assertEqual(self.database.get_prix_offre("91", 1, 1), 10.0)
        self.assertEqual(self.database.get_montant_total_ht("91", 1), 35.0)
        self.assertEqual(
            tuple(self.database.get_lot_montants_totaux("91", 1)),
            (35.0, 6.65, 41.65),
        )

    def test_partial_dqe_price_save_blocks_stale_totals(self):
        self.database.insert_soumissionnaire("92", "Brouillon", "1")
        self.database.add_offre_qualifie("92", 1, 85)
        self.database.replace_articles_from_dqe({1: [(1, 1), (2, 1)]})
        self.database.save_dqe_lot("92", 1, [(1, 10, 10), (2, 20, 20)])

        self.database.save_dqe_price_draft("92", 1, 1, 12, 12)

        self.assertEqual(self.database.get_prix_offre("92", 1, 1), 12.0)
        self.assertEqual(
            self.database.get_financial_readiness("92", 1),
            ["total DQE invalide", "montant de lettre non saisi"],
        )

    def test_changed_dqe_is_blocked_when_financial_values_exist(self):
        self.database.insert_soumissionnaire("9", "Test", "1")
        self.database.add_offre_qualifie("9", 1, 85)
        self.database.replace_articles_from_dqe({1: [(1, 10)]})
        self.database.insert_prix_offre("9", 1, 1, 5)
        self.database.insert_article_montants_ht("9", 1, 1, 50)
        self.database.insert_lot_montants_totaux("9", 1, 50, 9.5, 59.5)

        self.database.replace_articles_from_dqe({1: [(1, 10)]})
        with self.assertRaisesRegex(ValueError, "prix ou montants saisis"):
            self.database.replace_articles_from_dqe({1: [(1, 12)]})

        self.assertEqual(self.database.fetch_articles_by_lot(1), [(1, 10.0)])
        self.assertEqual(self.database.get_prix_offre("9", 1, 1), 5.0)
        self.assertEqual(self.database.get_lot_montants_totaux("9", 1)[0], 50.0)

    def test_dqe_import_records_checksum_and_invalid_rows(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "LOT 1"
        sheet.append(["Article", "Quantité"])
        sheet.append([1, 10])
        sheet.append(["bad", 4])
        path = tempfile.mktemp(suffix=".xlsx")
        workbook.save(path)
        try:
            result = import_dqe_workbook(self.database, path, return_details=True)
            self.assertEqual(result["total_articles"], 1)
            self.assertEqual(len(result["invalid_rows"]), 1)
            self.assertEqual(self.database.get_last_dqe_import()[3], 1)
            self.assertEqual(len(self.database.get_last_dqe_import()[0]), 64)
            self.assertTrue(os.path.exists(f"{self.database_path}.bak"))
        finally:
            os.remove(path)

    def test_financial_readiness_requires_dqe_and_letter_amount(self):
        self.assertEqual(self.database.get_financial_readiness("9", 1), ["DQE non enregistré"])
        self.database.insert_soumissionnaire("9", "Test", "1")
        self.database.add_offre_qualifie("9", 1, 85)
        self.assertEqual(
            [tuple(row) for row in self.database.get_soumissionnaires_qualifies()],
            [("9", "Test", "1")],
        )
        self.assertIsNone(
            self.database._schema_object_type('soumissionnaire_qualifie')
        )
        self.database.insert_lot_montants_totaux("9", 1, 100, 19, 119)
        self.assertEqual(
            self.database.get_financial_readiness("9", 1),
            ["montant de lettre non saisi"],
        )

    def test_classement_validation_persists_and_is_audited(self):
        self.assertFalse(self.database.is_classement_validated(5))
        self.database.validate_classement(5, "Test de validation")
        self.assertTrue(self.database.is_classement_validated(5))
        audit = self.database.get_audit_log(1)[0]
        self.assertEqual(audit[0], "valider")
        self.assertEqual(audit[2], "5")

    def test_technical_recalculation_removes_stale_qualification(self):
        self.database.insert_soumissionnaire("9", "Test", "1")
        self.database.insert_offre("9", 1, 1, 1, 6, 1, 1, 1, 2, 5, 60, 24)
        self.database.enregistrer_soumissionnaires_qualifies()
        self.assertEqual(self.database.get_notes_for_soumissionnaire("9")[0][0], 1)
        self.assertTrue(self.database.has_qualified_lots())

        self.database.update_offre(
            "9", 1,
            {
                "ing": 0, "ts": 0, "ouvrier": 0, "metreur": 0,
                "camion": 1, "nacelle": 1, "moyen_de_levage": 2,
                "outillage": 5, "delai_execution": 90, "delai_garantie": 24,
            },
        )
        self.database.enregistrer_soumissionnaires_qualifies()
        self.assertFalse(self.database.has_qualified_lots())
        self.assertEqual(self.database.get_notes_for_soumissionnaire("9"), [])
        self.assertEqual(
            self.database.get_offer_summary()[0][3],
            "éliminé",
        )

    def test_v2_migration_preserves_existing_technical_results(self):
        self.database.close()
        connection = sqlite3.connect(self.database_path)
        connection.execute('PRAGMA foreign_keys = OFF')
        connection.execute('DROP VIEW offre_qualifie')
        for table in (
            'lot_montants_totaux', 'article_montants_ht', 'prix_offre',
            'workflow_lot', 'articles', 'estimation', 'classement_validation',
        ):
            connection.execute(f'DROP TABLE {table}')
        connection.execute('''
            CREATE TABLE workflow_lot (
                n_pli TEXT NOT NULL,
                lot INTEGER NOT NULL,
                statut TEXT NOT NULL DEFAULT 'à vérifier',
                note_tech REAL,
                valide INTEGER NOT NULL DEFAULT 0,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                PRIMARY KEY (n_pli, lot),
                FOREIGN KEY(n_pli) REFERENCES soumissionnaire(n_pli)
            )
        ''')
        connection.execute('''
            CREATE TABLE articles (
                lot_num INTEGER,
                num_article INTEGER,
                quantite REAL,
                PRIMARY KEY (lot_num, num_article)
            )
        ''')
        connection.execute('''
            CREATE TABLE prix_offre (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                n_pli TEXT,
                lot_num INTEGER,
                num_article TEXT,
                prix_unitaire REAL,
                FOREIGN KEY(n_pli) REFERENCES soumissionnaire_qualifie(n_pli),
                FOREIGN KEY(lot_num, num_article) REFERENCES articles(lot_num, num_article),
                UNIQUE(n_pli, lot_num, num_article)
            )
        ''')
        connection.execute('''
            CREATE TABLE article_montants_ht (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                n_pli TEXT,
                lot_num INTEGER,
                num_article INTEGER,
                montant_ht REAL,
                FOREIGN KEY(n_pli) REFERENCES soumissionnaire_qualifie(n_pli),
                FOREIGN KEY(lot_num, num_article) REFERENCES articles(lot_num, num_article),
                UNIQUE(n_pli, lot_num, num_article)
            )
        ''')
        connection.execute('''
            CREATE TABLE lot_montants_totaux (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                n_pli TEXT,
                lot_num INTEGER,
                montant_total_ht REAL,
                montant_total_tva REAL,
                montant_total_ttc REAL,
                montant_lettre_soumission_ttc REAL,
                taux REAL,
                FOREIGN KEY(n_pli) REFERENCES soumissionnaire_qualifie(n_pli),
                FOREIGN KEY(lot_num) REFERENCES articles(lot_num),
                UNIQUE(n_pli, lot_num)
            )
        ''')
        connection.execute('''
            CREATE TABLE estimation (
                lot_num INTEGER,
                montant_lettre_soumission REAL,
                PRIMARY KEY (lot_num, montant_lettre_soumission)
            )
        ''')
        connection.execute('''
            CREATE TABLE classement_validation (
                lot_num INTEGER PRIMARY KEY,
                validated_by TEXT NOT NULL DEFAULT 'utilisateur local',
                validated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                details TEXT
            )
        ''')
        connection.execute('''
            CREATE TABLE soumissionnaire_qualifie (
                n_pli TEXT PRIMARY KEY,
                nom TEXT,
                lots TEXT,
                statut TEXT NOT NULL DEFAULT 'qualifié'
            )
        ''')
        connection.execute(
            "INSERT INTO soumissionnaire (n_pli, nom, lots) VALUES ('10', 'Legacy', '1,2,3')"
        )
        connection.execute('''
            INSERT INTO offre (
                n_pli, lot, ing, ts, ouvrier, metreur, camion, nacelle,
                moyen_de_levage, outillage, delai_execution, delai_garantie
            ) VALUES ('99', 4, 1, 0, 2, 1, 1, 0, 1, 2, 60, 12)
        ''')
        connection.execute('''
            INSERT INTO workflow_lot (n_pli, lot, statut, note_tech, valide)
            VALUES ('10', 1, 'éliminé', 42, 0)
        ''')
        connection.execute('''
            CREATE TABLE offre_qualifie (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                n_pli TEXT, lot INTEGER, noteTech REAL
            )
        ''')
        connection.executemany(
            'INSERT INTO offre_qualifie (n_pli, lot, noteTech) VALUES (?, ?, ?)',
            [('10', 1, 98), ('10', 2, 85)],
        )
        connection.execute('''
            INSERT INTO soumissionnaire_qualifie (n_pli, nom, lots, statut)
            VALUES ('10', 'Legacy', '2,3', 'qualifié')
        ''')
        connection.execute(
            'INSERT INTO articles (lot_num, num_article, quantite) VALUES (1, 1, 10)'
        )
        connection.execute('''
            INSERT INTO prix_offre (n_pli, lot_num, num_article, prix_unitaire)
            VALUES ('10', 1, '1', 5)
        ''')
        connection.execute('''
            INSERT INTO article_montants_ht (n_pli, lot_num, num_article, montant_ht)
            VALUES ('10', 1, 1, 50)
        ''')
        connection.execute('''
            INSERT INTO lot_montants_totaux (
                n_pli, lot_num, montant_total_ht, montant_total_tva, montant_total_ttc
            ) VALUES ('10', 1, 50, 9.5, 59.5)
        ''')
        connection.execute('PRAGMA user_version = 1')
        connection.commit()
        connection.close()

        self.database = Database(self.database_path)
        rows = self.database.cursor.execute('''
            SELECT lot, statut, note_tech FROM workflow_lot
            WHERE n_pli = '10' ORDER BY lot
        ''').fetchall()
        self.assertEqual(
            [tuple(row) for row in rows],
            [(1, 'éliminé', 42.0), (2, 'qualifié', 85.0), (3, 'qualifié', None)],
        )
        self.assertEqual(self.database.get_prix_offre('10', 1, 1), 5.0)
        self.assertEqual(self.database.get_montant_total_ht('10', 1), 50.0)
        self.assertEqual(
            tuple(self.database.get_lot_montants_totaux('10', 1)),
            (50.0, 9.5, 59.5),
        )
        self.assertEqual(self.database.get_schema_version(), 4)
        self.assertEqual(
            self.database._schema_object_type('offre_qualifie'),
            'view',
        )
        self.assertIsNone(
            self.database._schema_object_type('soumissionnaire_qualifie')
        )
        recovered_bidder = self.database.cursor.execute(
            "SELECT nom FROM soumissionnaire WHERE n_pli = '99'"
        ).fetchone()[0]
        self.assertEqual(recovered_bidder, 'Fiche récupérée - Pli 99')
        self.assertIsNotNone(self.database.get_offre('99', 4))
        recovery_log = self.database.cursor.execute(
            "SELECT action FROM audit_log WHERE entity_id = '99'"
        ).fetchone()
        self.assertEqual(recovery_log[0], 'recuperer')
        self.assertTrue(os.path.exists(f'{self.database_path}.bak'))

    def test_financial_foreign_keys_are_enabled_and_consistent(self):
        self.assertEqual(
            self.database.cursor.execute('PRAGMA foreign_keys').fetchone()[0],
            1,
        )
        self.assertEqual(
            self.database.cursor.execute('PRAGMA foreign_key_check').fetchall(),
            [],
        )
        workflow_foreign_keys = self.database.cursor.execute(
            'PRAGMA foreign_key_list(prix_offre)'
        ).fetchall()
        self.assertIn('workflow_lot', [row[2] for row in workflow_foreign_keys])
        self.database.replace_articles_from_dqe({1: [(1, 10)]})
        with self.assertRaises(sqlite3.IntegrityError):
            self.database.insert_prix_offre('missing-bidder', 1, 1, 5)
        self.database.conn.rollback()
        self.assertIsNone(self.database.get_prix_offre('missing-bidder', 1, 1))

    def test_offer_mutations_are_centralized_and_audited(self):
        self.database.insert_soumissionnaire("9", "Test", "1")
        self.database.insert_offre("9", 1, 1, 0, 2, 1, 1, 0, 1, 2, 60, 12)
        self.database.update_offre(
            "9", 1,
            {
                "ing": 0, "ts": 1, "ouvrier": 3, "metreur": 1,
                "camion": 0, "nacelle": 1, "moyen_de_levage": 2,
                "outillage": 4, "delai_execution": 90, "delai_garantie": 24,
            },
        )
        self.assertEqual(self.database.get_offre("9", 1)[3:13], (0, 1, 3, 1, 0, 1, 2, 4, 90, 24))
        audit = self.database.get_audit_log(1)[0]
        self.assertEqual(audit[0], "modifier")
        self.assertIn('"ing": 1', audit[6])
        self.assertIn('"ing": 0', audit[7])

        self.database.delete_offre("9", 1)
        self.assertIsNone(self.database.get_offre("9", 1))
        self.assertEqual(self.database.get_audit_log(1)[0][0], "supprimer")

    def test_validated_lot_rejects_financial_changes_until_reopened(self):
        self.database.insert_soumissionnaire("9", "Test", "1")
        self.database.add_offre_qualifie("9", 1, 85)
        self.database.replace_articles_from_dqe({1: [(1, 2)]})
        self.database.insert_prix_offre("9", 1, 1, 10)
        self.database.insert_lot_montants_totaux("9", 1, 20, 3.8, 23.8)
        self.database.validate_classement(1)

        with self.assertRaises(ValueError):
            self.database.insert_prix_offre("9", 1, 1, 12)
        with self.assertRaises(ValueError):
            self.database.update_montant_lettre("9", 1, 25)

        self.database.reopen_classement(1, "Correction demandée")
        self.database.insert_prix_offre("9", 1, 1, 12)
        self.assertFalse(self.database.is_classement_validated(1))
        self.assertEqual(self.database.get_prix_offre("9", 1, 1), 12)

    def test_dashboard_summary_exposes_actionable_indicators(self):
        summary = self.database.get_dashboard_summary()
        self.assertEqual(summary["soumissionnaires"], 0)
        self.assertEqual(summary["offres_incompletes"], 0)
        self.assertEqual(summary["lots_bloques"], 0)
        self.assertIsNone(summary["derniere_activite"])

    def test_negative_financial_values_are_rejected(self):
        with self.assertRaises(ValueError):
            self.database.insert_prix_offre("9", 1, 1, -1)
        with self.assertRaises(ValueError):
            self.database.insert_lot_montants_totaux("9", 1, -10, 0, 0)

    def test_ui_theme_scales_decimal_font_sizes(self):
        from ui_theme import _scaled_stylesheet

        stylesheet = _scaled_stylesheet(1.2)
        self.assertIn("font-size: 13px", stylesheet)

    def test_word_export_layout_is_landscape_and_fits_printable_width(self):
        from io import BytesIO

        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.oxml.ns import qn
        from docx.shared import Inches

        from word_export import configure_landscape_document, format_landscape_table

        document = Document()
        usable_width = configure_landscape_document(document)
        table = document.add_table(rows=2, cols=3)
        for index, text in enumerate(("Pli", "Soumissionnaire", "Montant TTC")):
            table.cell(0, index).text = text
        for index, text in enumerate(("18", "Ahmed", "1 234 567,00")):
            table.cell(1, index).text = text
        format_landscape_table(table, (0.8, 1.8, 1.2), usable_width)

        output = BytesIO()
        document.save(output)
        output.seek(0)
        exported = Document(output)
        section = exported.sections[0]
        exported_table = exported.tables[0]
        table_width_twips = int(
            exported_table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w"))
        )
        header_properties = exported_table.rows[0]._tr.get_or_add_trPr()

        self.assertEqual(section.orientation, WD_ORIENT.LANDSCAPE)
        self.assertGreater(section.page_width, section.page_height)
        self.assertEqual(
            table_width_twips,
            round(usable_width / Inches(1) * 1440),
        )
        self.assertIsNotNone(header_properties.find(qn("w:tblHeader")))

    def test_word_export_formatting_is_idempotent_when_applied_more_than_once(self):
        from docx import Document
        from docx.oxml.ns import qn

        from word_export import configure_landscape_document, format_landscape_table

        document = Document()
        usable_width = configure_landscape_document(document)
        table = document.add_table(rows=3, cols=3)
        table.cell(0, 0).text = "Pli"
        table.cell(0, 1).text = "Soumissionnaire"
        table.cell(0, 2).text = "Montant"
        table.cell(1, 0).text = "18"
        table.cell(1, 1).text = "Ahmed"
        table.cell(1, 2).text = "1 234 567,00"

        format_landscape_table(table, (0.8, 1.8, 1.2), usable_width)
        format_landscape_table(table, (0.8, 1.8, 1.2), usable_width)

        header_properties = table.rows[0]._tr.get_or_add_trPr()
        first_data_row_properties = table.rows[1]._tr.get_or_add_trPr()

        self.assertEqual(len(header_properties.findall(qn("w:tblHeader"))), 1)
        self.assertEqual(len(first_data_row_properties.findall(qn("w:cantSplit"))), 1)

    def test_soumissionnaire_search_filters_by_text_and_lot(self):
        self.database.insert_soumissionnaire("10", "Alpha", "1")
        self.database.insert_soumissionnaire("20", "Beta", "2")
        self.database.insert_offre("10", 1, 1, 0, 1, 1, 1, 0, 0, 1, 60, 12)
        self.database.insert_offre("20", 2, 1, 0, 1, 1, 1, 0, 0, 1, 60, 12)
        self.assertEqual(len(self.database.search_soumissionnaires("Alpha")), 1)
        self.assertEqual(self.database.search_soumissionnaires("", 2)[0][0], "20")

    def test_registry_search_matches_exact_bidder_number_and_combines_filters(self):
        for n_pli, nom in (("1", "Ahmed Ali"), ("10", "Ahmed Naji"), ("11", "Ahmed Test")):
            self.database.insert_soumissionnaire(n_pli, nom, "1 2")
        self.database.insert_offre("1", 1, 1, 0, 1, 1, 1, 0, 0, 1, 60, 12)
        self.database.insert_offre("1", 2, 1, 0, 1, 1, 1, 0, 0, 1, 60, 12)
        self.database.insert_offre("10", 1, 1, 0, 1, 1, 1, 0, 0, 1, 60, 12)
        self.database.insert_offre("11", 2, 1, 0, 1, 1, 1, 0, 0, 1, 60, 12)

        exact_match = self.database.search_registry_soumissionnaires(n_pli="1")
        self.assertEqual([row[0] for row in exact_match], ["1"])

        name_and_lot = self.database.search_registry_soumissionnaires(
            nom="ahmed", lot=1
        )
        self.assertEqual([row[0] for row in name_and_lot], ["1", "10"])

        combined = self.database.search_registry_soumissionnaires(
            n_pli="1", nom="ali", lot=2
        )
        self.assertEqual([row[0] for row in combined], ["1"])

    def test_offer_summary_contains_bidder_lot_and_status(self):
        self.database.insert_soumissionnaire("10", "Alpha", "1")
        self.database.insert_offre("10", 1, 1, 0, 1, 1, 1, 0, 0, 1, 60, 12)
        summary = self.database.get_offer_summary()
        self.assertEqual(summary, [("10", "Alpha", 1, "à vérifier")])

    def test_offer_indexes_cover_lot_and_bidder_lookups(self):
        lot_plan = self.database.cursor.execute(
            "EXPLAIN QUERY PLAN SELECT COUNT(*) FROM offre WHERE lot = ?",
            (1,),
        ).fetchall()
        bidder_plan = self.database.cursor.execute(
            "EXPLAIN QUERY PLAN SELECT COUNT(*) FROM offre WHERE n_pli = ? AND lot = ?",
            ("10", 1),
        ).fetchall()
        self.assertIn("idx_offre_lot_npli", " ".join(row[3] for row in lot_plan))
        self.assertIn("idx_offre_npli_lot", " ".join(row[3] for row in bidder_plan))

        self.database.cursor.execute("DROP INDEX idx_offre_lot_npli")
        self.database.cursor.execute("DROP INDEX idx_offre_npli_lot")
        self.database.conn.commit()
        migrated_database = Database(self.database_path)
        try:
            migrated_plan = migrated_database.cursor.execute(
                "EXPLAIN QUERY PLAN SELECT COUNT(*) FROM offre WHERE lot = ?",
                (1,),
            ).fetchall()
            self.assertIn(
                "idx_offre_lot_npli",
                " ".join(row[3] for row in migrated_plan),
            )
        finally:
            migrated_database.close()

    def test_shared_integer_validation_preserves_workflow_limits(self):
        self.assertEqual(parse_integer("11", minimum=1, maximum=11), 11)
        self.assertIsNone(parse_integer("12", minimum=1, maximum=11))
        self.assertIsNone(parse_integer("lot 1", minimum=1, maximum=11))
        self.assertTrue(is_digit_string("0012"))
        self.assertEqual(parse_integer_tokens("1 invalide 3"), [1, 3])
        self.assertEqual(parse_integer("45", minimum=0), 45)
        self.assertIsNone(parse_integer("-1", minimum=0))

    def test_shared_amount_validation_accepts_local_decimal_input(self):
        self.assertEqual(parse_amount("1 234,50", maximum_decimal_places=2), 1234.5)
        self.assertEqual(parse_amount("12,5000", maximum_decimal_places=2), 12.5)
        self.assertIsNone(parse_amount("12,345", maximum_decimal_places=2))
        self.assertIsNone(parse_amount("-1"))
        self.assertIsNone(parse_amount("NaN"))

    def test_lot_delays_respect_minimum_requirements(self):
        from input_validation import validate_lot_delay_inputs

        self.assertIsNone(validate_lot_delay_inputs(1, 2, 2))
        self.assertIsNone(validate_lot_delay_inputs(3, 2, 2))
        self.assertEqual(validate_lot_delay_inputs(1, 60, 12), (60, 12))
        self.assertEqual(validate_lot_delay_inputs(3, 45, 12), (45, 12))

    def test_technical_summary_keeps_identity_score_and_status(self):
        row = ["10", "Alpha"] + [0] * 12 + [82.5, "qualifié"]
        self.assertEqual(
            TechnicalEvaluation.build_summary_rows([row]),
            [["10", "Alpha", 82.5, "qualifié"]],
        )


if __name__ == "__main__":
    unittest.main()
